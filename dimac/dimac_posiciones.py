#!/usr/bin/env python3
"""Baja TODAS las posiciones de Dimac (N° Dimac) a un Excel para cargar en el ERP.

Credenciales por variables de entorno (igual que dimac_buscar.py):
    DIMAC_USUARIO, DIMAC_CLAVE

Uso:
    python dimac/dimac_posiciones.py
    python dimac/dimac_posiciones.py --salida salida/dimac_posiciones.xlsx

Hojas del Excel:
  Posiciones    una fila por N° Dimac: tipo, descripción, cada medida en su
                columna, vehículos, marcas que la fabrican, imagen y URL.
  Aplicaciones  una fila por posición y vehículo: marca, modelo, años y eje.
  Articulos     una fila por artículo (marca + calidad) con su precio de lista.

Fuentes: la descarga del catálogo completo (búsqueda vacía, view=download),
la misma descarga filtrada por "Tipo Articulo" para clasificar cada posición, y
las páginas de resultados (10 por página) que traen los vehículos ya separados,
la imagen y la URL.
"""
import argparse
import os
import re
import sys
import time
from collections import defaultdict
from pathlib import Path

import requests
from bs4 import BeautifulSoup

from dimac_buscar import (
    BASE, UA, limpiar, login, params_busqueda, parsear, pedir,
    parsear_descarga, separar_aplicaciones,
)

# Orden de las columnas de medidas; las que no estén acá se agregan al final.
ORDEN_MEDIDAS = [
    "Sistema Freno", "Largo Interno", "Alto Interno", "Espesor Interno",
    "Largo Externo", "Alto Externo", "Espesor Externo",
    "Largo", "Ancho", "Diametro", "Espesor Sub-STD", "Espesor STD",
    "Espesor 1era", "Espesor 2da", "Espesor 3era", "Espesor 4ta",
    "Cant.de Agujeros", "Diametro de Agujero", "Cant. de Piezas por Juego",
    "Anclajes de Fijacion", "Anclajes de Fijacion Adicional",
    "Sensor Mecanico", "Sensor Electrico", "Detalle del Sensor y Cant.",
]
EJES = r"delantera/trasera|delantera|trasera|tercer eje"


def filtros(html, titulo):
    """Valores de un filtro de la búsqueda: [(texto, valor para 'filters')]."""
    soup = BeautifulSoup(html, "html.parser")
    for t in soup.select(".filter_title"):
        if titulo.lower() in t.get_text().lower():
            area = t.find_next(class_="filter_area")
            return [
                (limpiar(a.get_text()), re.search(r"filters=([^&]*)", a["href"]).group(1))
                for a in area.select("a.filter_elem[href]")
            ]
    return []


def numero(valor):
    """'139.0' -> 139.0; deja el texto si no es un número puro."""
    return float(valor) if re.fullmatch(r"\d+(\.\d+)?", valor) else valor


def medidas_a_dict(texto):
    datos = {}
    for parte in texto.split(";"):
        if ":" in parte:
            clave, valor = parte.split(":", 1)
            datos[clave.strip()] = numero(valor.strip())
    return datos


def anio(txt):
    if not txt:
        return None
    n = int(txt)
    if len(txt) == 2:
        n += 1900 if n >= 30 else 2000
    return n


def separar_vehiculo(texto, marcas):
    """'FORD ---Focus II 1.6 Turnier 04/ trasera' -> marca, modelo, desde, hasta, eje."""
    t = limpiar(texto)
    marca = next((m for m in marcas if t.upper().startswith(m.upper() + " ") or t.upper() == m.upper()), "")
    resto = t[len(marca):].strip() if marca else t
    eje = ""
    m = re.search(rf"\s*\b({EJES})\s*$", resto, re.I)
    if m:
        eje, resto = m.group(1).lower(), resto[:m.start()]
    desde = hasta = None
    # "90/94", "04/", "2017/", "/ 96"; no "E24/27" ni "1200/1300".
    m = re.search(r"(?<![\w./])(\d{2}|\d{4})?\s*/\s*(\d{2}|\d{4})?\s*$", resto)
    if m and (m.group(1) or m.group(2)):
        d, h = anio(m.group(1)), anio(m.group(2))
        if all(a is None or 1930 <= a <= 2035 for a in (d, h)) and not (d and h and h < d):
            desde, hasta = d, h
            resto = resto[:m.start()]
    modelo = limpiar(resto).lstrip("-").strip()
    return {"marca_vehiculo": marca, "modelo": modelo, "anio_desde": desde, "anio_hasta": hasta, "eje": eje}


def separar_articulo(codigo, linea):
    """'of 930550', 'PREMIX  Zapatas 1ra' -> código, oferta, fabricante, producto, calidad."""
    oferta = codigo.lower().startswith("of ")
    codigo = codigo[3:].strip() if oferta else codigo
    linea = limpiar(linea)
    m = re.match(r"(.*?)\s*\b(Pastillas|Zapatas|Cintas Perforadas|Cintas Ciegas|Cintas|Placa)\b\s*(.*)", linea)
    if m and m.group(1):
        fabricante, producto, calidad = m.group(1), m.group(2), m.group(3)
    else:  # "Pastillas en Desarrollo", "Placa Antirruido"
        fabricante, producto, calidad = "", linea, ""
    return {
        "codigo_articulo": codigo, "fabricante": fabricante, "producto": producto,
        "calidad": calidad, "oferta": "Sí" if oferta else "", "linea_dimac": linea,
    }


def guardar(hojas, destino):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    anchos = {"descripcion": 45, "aplicaciones": 80, "texto": 60, "modelo": 40,
              "marcas_vehiculo": 30, "fabricantes": 35, "url": 50, "imagen": 45, "linea_dimac": 35}
    for nombre, (columnas, filas) in hojas.items():
        ws = wb.create_sheet(nombre)
        ws.append(columnas)
        for f in filas:
            ws.append([f.get(c) for c in columnas])
        for i, col in enumerate(columnas, start=1):
            ws.cell(1, i).font = Font(bold=True)
            ws.column_dimensions[ws.cell(1, i).column_letter].width = anchos.get(col, max(12, min(len(col) + 2, 28)))
            if col == "precio_lista":
                for (celda,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    celda.number_format = "#,##0.00"
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(destino)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--salida", default="salida/dimac_posiciones.xlsx")
    ap.add_argument("--max-paginas", type=int, default=400)
    args = ap.parse_args()

    usuario = os.environ.get("DIMAC_USUARIO")
    clave = os.environ.get("DIMAC_CLAVE")
    if not usuario or not clave:
        sys.exit("Faltan las variables de entorno DIMAC_USUARIO y DIMAC_CLAVE.")

    sesion = requests.Session()
    sesion.headers["User-Agent"] = UA
    login(sesion, usuario, clave)
    print("Login OK")

    def descarga(**extra):
        r = pedir(sesion, BASE + "/search", params_busqueda("", view="download", **extra))
        return parsear_descarga(r.content.decode("utf-8", errors="replace"))

    # 1) Catálogo completo: una fila por artículo.
    articulos = descarga()
    print(f"Catálogo: {len(articulos)} artículos")

    # 2) Página 1: filtros (tipos y marcas de vehículo) + primeras posiciones.
    html = pedir(sesion, BASE + "/search", params_busqueda("", 1)).text
    marcas = sorted((t for t, _ in filtros(html, "Marca Vehiculo")), key=len, reverse=True)
    tipo_de = {}
    for tipo, valor in filtros(html, "Tipo Articulo"):
        for a in descarga(filters=valor):
            tipo_de[a["codigo_dimac"].upper()] = tipo
    print(f"Tipos asignados a {len(tipo_de)} posiciones")

    # 3) Páginas de resultados: URL, imagen y vehículos separados.
    web = {}
    pagina = 1
    while True:
        soup = BeautifulSoup(html, "html.parser")
        productos, hay_siguiente = parsear(html)
        web.update(productos)
        for caja in soup.select("article.PRODUCT_BOX"):
            codigo = limpiar(caja.select_one("h3").get_text()).replace("DIMAC NRO.", "").strip().upper()
            mme = caja.select_one(".mme")
            web[codigo]["vehiculos"] = [limpiar(x) for x in mme.stripped_strings] if mme else []
        if not hay_siguiente or pagina >= args.max_paginas:
            break
        pagina += 1
        if pagina % 25 == 0:
            print(f"  página {pagina}")
        time.sleep(0.2)
        html = pedir(sesion, BASE + "/search", params_busqueda("", pagina)).text
    print(f"Páginas leídas: {pagina} ({len(web)} posiciones con URL e imagen)")

    # 4) Armar las hojas.
    por_posicion = defaultdict(list)
    for a in articulos:
        por_posicion[a["codigo_dimac"].upper()].append(a)

    hoja_pos, hoja_apl, hoja_art = [], [], []
    claves_medidas = []
    for codigo in sorted(por_posicion):
        arts = por_posicion[codigo]
        base = arts[0]
        info = web.get(codigo, {})
        vehiculos = info.get("vehiculos") or [v for v in separar_aplicaciones(base["aplicacion"]).split(" | ") if v]
        apls = [dict(separar_vehiculo(v, marcas), posicion=base["codigo_dimac"], texto=v) for v in vehiculos]
        hoja_apl.extend(apls)

        filas_art = [dict(separar_articulo(a["codigo_articulo"], a["marca"]),
                          posicion=base["codigo_dimac"], precio_lista=a["precio_lista"]) for a in arts]
        hoja_art.extend(filas_art)

        medidas = medidas_a_dict(base["medidas"])
        claves_medidas.extend(k for k in medidas if k not in claves_medidas)
        hoja_pos.append({
            "posicion": base["codigo_dimac"],
            "tipo": tipo_de.get(codigo, ""),
            "descripcion": base["descripcion"].rstrip("- ").strip(),
            "eje": " / ".join(sorted({a["eje"] for a in apls if a["eje"]})),
            **medidas,
            "cant_vehiculos": len(apls),
            "marcas_vehiculo": ", ".join(sorted({a["marca_vehiculo"] for a in apls if a["marca_vehiculo"]})),
            "aplicaciones": " | ".join(vehiculos),
            "fabricantes": ", ".join(sorted({f["fabricante"] or f["producto"] for f in filas_art})),
            "cant_articulos": len(arts),
            "imagen": info.get("imagen", ""),
            "url": info.get("url", ""),
        })

    medidas_cols = [k for k in ORDEN_MEDIDAS if k in claves_medidas] + [k for k in claves_medidas if k not in ORDEN_MEDIDAS]
    hojas = {
        "Posiciones": (["posicion", "tipo", "descripcion", "eje", *medidas_cols, "cant_vehiculos",
                        "marcas_vehiculo", "aplicaciones", "fabricantes", "cant_articulos", "imagen", "url"], hoja_pos),
        "Aplicaciones": (["posicion", "marca_vehiculo", "modelo", "anio_desde", "anio_hasta", "eje", "texto"], hoja_apl),
        "Articulos": (["posicion", "codigo_articulo", "fabricante", "producto", "calidad", "oferta",
                       "precio_lista", "linea_dimac"], hoja_art),
    }
    destino = Path(args.salida)
    guardar(hojas, destino)
    print(f"{len(hoja_pos)} posiciones, {len(hoja_apl)} aplicaciones y {len(hoja_art)} artículos guardados en {destino}")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Inicia sesión en dimac.com.ar y busca productos (ej. pastillas de freno).

Credenciales por variables de entorno (nunca en el código):
    DIMAC_USUARIO   usuario (CUIT sin guiones)
    DIMAC_CLAVE     contraseña

Uso:
    python dimac/dimac_buscar.py "pastilla"
    python dimac/dimac_buscar.py "gol trend"
    python dimac/dimac_buscar.py "" --sin-stock          # catálogo completo
    python dimac/dimac_buscar.py "corsa" --salida salida/corsa.xlsx

De dónde sale cada dato:
  - "Descargar esta búsqueda" (/search?...&view=download): una fila por
    artículo con N° Dimac, fabricante, código, precio de lista, descripción,
    aplicaciones (marca-modelo-eje) y medidas. Es la fuente principal.
  - Páginas de resultados (10 por página, fijo): URL e imagen de cada N° Dimac.
  - Ficha de cada producto: stock de cada artículo (una visita por producto,
    por eso se limita con --max-fichas).

Guarda el Excel y el HTML crudo en salida/html/ para revisar si el sitio cambia.
"""
import argparse
import csv
import os
import re
import sys
import time
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup

BASE = "https://www.dimac.com.ar"
UA = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
COLUMNAS = [
    "codigo_dimac", "codigo_articulo", "marca", "descripcion", "aplicacion",
    "medidas", "precio_lista", "stock", "url", "imagen",
]


def login(sesion, usuario, clave):
    sesion.get(BASE + "/", timeout=30)
    r = sesion.post(
        BASE + "/User/LoginBox",
        data={"Username": usuario, "Password": clave},
        headers={"Referer": BASE + "/"},
        timeout=30,
    )
    r.raise_for_status()
    # La respuesta del POST siempre trae el formulario de ingreso, aun con login
    # correcto; lo que indica éxito es la cookie de sesión .ASPXAUTH.
    if ".ASPXAUTH" not in sesion.cookies:
        sys.exit("Login fallido: revisá DIMAC_USUARIO (CUIT sin guiones) y DIMAC_CLAVE.")


def params_busqueda(termino, pagina=1, **extra):
    p = {
        "page": pagina,
        "id": 0,
        "recsPerPage": 10,  # el sitio ignora valores mayores
        "order": "Code",
        "sort": "False",
        "itemtype": "Product",
        "term": termino,
        "getFilterData": "True",
    }
    p.update(extra)
    return p


def pedir(sesion, url, params=None):
    r = sesion.get(url, params=params, timeout=180)
    if r.status_code == 500 and "CurrentUser" in r.text:
        sys.exit("El sitio no reconoce la sesión (búsqueda requiere login).")
    r.raise_for_status()
    return r


def limpiar(texto):
    return re.sub(r"\s+", " ", texto or "").strip()


def precio_a_numero(texto):
    t = re.sub(r"[^\d,.-]", "", texto or "").replace(".", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def separar_aplicaciones(texto):
    # El sitio pega las aplicaciones sin separador:
    # "DODGE - CHRYSLER Sebring 03/05 traseraDODGE - CHRYSLER Stratus 03/06 trasera"
    partes = re.split(r"(?<=[a-z0-9).])(?=[A-Z]{2})", limpiar(texto))
    return " | ".join(p.strip() for p in partes if p.strip())


def parsear_descarga(html):
    """Tabla de "Descargar esta búsqueda": una fila por artículo."""
    soup = BeautifulSoup(html, "html.parser")
    filas = []
    for tr in soup.select("tr"):
        celdas = tr.find_all("td")
        if len(celdas) < 7:
            continue
        txt = [c.get_text() for c in celdas]
        filas.append({
            "codigo_dimac": limpiar(txt[0]).strip('="'),
            "codigo_articulo": limpiar(txt[2]),
            "marca": limpiar(txt[1]),
            "descripcion": limpiar(txt[4]),
            "aplicacion": separar_aplicaciones(txt[5]),
            "medidas": "; ".join(limpiar(x) for x in txt[6].splitlines() if x.strip()),
            "precio_lista": precio_a_numero(txt[3]),
            "stock": "",
            "url": "",
            "imagen": "",
        })
    return filas


def parsear(html):
    """Página de resultados: devuelve ({codigo_dimac: {url, imagen}}, hay_siguiente)."""
    soup = BeautifulSoup(html, "html.parser")
    productos = {}
    for caja in soup.select("article.PRODUCT_BOX"):
        titulo = caja.select_one("h3")
        enlace = caja.select_one("a.onbody[href]")
        img = caja.select_one(".box_image img[src]")
        if not titulo:
            continue
        codigo = limpiar(titulo.get_text()).replace("DIMAC NRO.", "").strip()
        productos[codigo.upper()] = {
            "url": urljoin(BASE, enlace["href"]) if enlace else "",
            "imagen": urljoin(BASE, img["src"]) if img else "",
        }
    return productos, soup.select_one(".pagination_wrapper a.p-next") is not None


def parsear_ficha(html):
    """Ficha de producto: {codigo_articulo: stock}."""
    soup = BeautifulSoup(html, "html.parser")
    stock = {}
    for tr in soup.select("table.table-r tbody tr"):
        codigo = tr.select_one(".codeid")
        marca = tr.select_one(".stock[data-title]")
        if codigo:
            # La ficha muestra "951022/2" donde la descarga dice "951022*2".
            stock[limpiar(codigo.get_text()).replace("/", "*")] = marca["data-title"] if marca else ""
    return stock


def guardar(filas, destino):
    destino.parent.mkdir(parents=True, exist_ok=True)
    if destino.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Dimac"
        ws.append(COLUMNAS)
        for f in filas:
            ws.append([f.get(c, "") for c in COLUMNAS])
        for celda in ws[1]:
            celda.font = Font(bold=True)
        anchos = {"descripcion": 45, "aplicacion": 70, "medidas": 35, "url": 50, "imagen": 40}
        for i, col in enumerate(COLUMNAS, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = anchos.get(col, 16)
        col_precio = COLUMNAS.index("precio_lista") + 1
        for fila in ws.iter_rows(min_row=2, min_col=col_precio, max_col=col_precio):
            fila[0].number_format = "#,##0.00"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(destino)
    else:
        with open(destino, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=COLUMNAS, delimiter=";")
            w.writeheader()
            w.writerows(filas)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("terminos", nargs="*", help='texto a buscar (vacío o "" = catálogo completo)')
    ap.add_argument("--max-paginas", type=int, default=200, help="páginas de 10 para sacar URL e imagen")
    ap.add_argument("--max-fichas", type=int, default=100, help="productos máximos a abrir para ver stock")
    ap.add_argument("--sin-stock", action="store_true", help="no abrir las fichas (más rápido)")
    ap.add_argument("--salida", default="salida/dimac_resultados.xlsx")
    args = ap.parse_args()

    usuario = os.environ.get("DIMAC_USUARIO")
    clave = os.environ.get("DIMAC_CLAVE")
    if not usuario or not clave:
        sys.exit("Faltan las variables de entorno DIMAC_USUARIO y DIMAC_CLAVE.")

    termino = " ".join(args.terminos).strip()
    sesion = requests.Session()
    sesion.headers["User-Agent"] = UA
    login(sesion, usuario, clave)
    print("Login OK")

    dir_html = Path("salida/html")
    dir_html.mkdir(parents=True, exist_ok=True)
    slug = re.sub(r"\W+", "_", termino).strip("_") or "catalogo"

    # 1) Datos principales: la descarga de la búsqueda.
    r = pedir(sesion, BASE + "/search", params_busqueda(termino, view="download"))
    html = r.content.decode("utf-8", errors="replace")
    (dir_html / f"{slug}_descarga.html").write_text(html, encoding="utf-8")
    filas = parsear_descarga(html)
    if not filas:
        print(f"No se encontraron productos para '{termino}'. Revisá el HTML en {dir_html}/.")
        return
    codigos = {f["codigo_dimac"].upper() for f in filas}
    print(f"{len(filas)} artículos ({len(codigos)} productos Dimac)")

    # 2) URL e imagen: páginas de resultados.
    enlaces = {}
    for pagina in range(1, args.max_paginas + 1):
        html = pedir(sesion, BASE + "/search", params_busqueda(termino, pagina)).text
        (dir_html / f"{slug}_p{pagina}.html").write_text(html, encoding="utf-8")
        productos, hay_siguiente = parsear(html)
        enlaces.update(productos)
        if not hay_siguiente or codigos <= enlaces.keys():
            break
    print(f"URL e imagen de {len(codigos & enlaces.keys())} productos")
    for f in filas:
        f.update(enlaces.get(f["codigo_dimac"].upper(), {}))

    # 3) Stock: ficha de cada producto.
    if args.sin_stock:
        pass
    elif len(codigos) > args.max_fichas:
        print(f"Stock no consultado: son {len(codigos)} productos (límite --max-fichas {args.max_fichas}).")
    else:
        stock = {}
        urls = sorted({f["url"] for f in filas if f["url"]})
        for i, url in enumerate(urls, 1):
            stock.update(parsear_ficha(pedir(sesion, url).text))
            if i % 20 == 0:
                print(f"  fichas {i}/{len(urls)}")
            time.sleep(0.3)
        for f in filas:
            f["stock"] = stock.get(f["codigo_articulo"], "")
        print(f"Stock consultado en {len(urls)} fichas")

    destino = Path(args.salida)
    guardar(filas, destino)
    print(f"{len(filas)} artículos guardados en {destino}")


if __name__ == "__main__":
    main()
